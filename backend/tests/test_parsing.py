"""Deterministic format parser tests."""

from io import BytesIO

import pytest
from docx import Document as DocxDocument

from thoughtharbor.documents.parsing import (
    ParserError,
    UnsupportedFormatError,
    parse_content,
)


def test_plain_text_normalizes_line_endings_and_preserves_offsets() -> None:
    parsed = parse_content(
        BytesIO(b"First line\r\n\r\nSecond line"),
        original_name="notes.md",
        media_type="text/markdown",
    )

    assert parsed.text == "First line\n\nSecond line"
    assert [section.text for section in parsed.sections] == ["First line", "Second line"]
    assert (
        parsed.text[parsed.sections[1].offset_start : parsed.sections[1].offset_end]
        == "Second line"
    )
    assert parsed.parser_name == "plain_text"


def test_email_parser_keeps_headers_as_metadata() -> None:
    raw = (
        b"From: sender@example.com\n"
        b"To: recipient@example.com\n"
        b"Subject: Project update\n"
        b"Date: Tue, 09 Sep 2025 10:00:00 +0000\n"
        b"Content-Type: text/plain; charset=utf-8\n\n"
        b"The body."
    )

    parsed = parse_content(BytesIO(raw), original_name="mail.eml", media_type="message/rfc822")

    assert parsed.text == "The body."
    assert parsed.metadata["subject"] == "Project update"
    assert parsed.metadata["sender"] == "sender@example.com"
    assert parsed.metadata["recipients"] == ["recipient@example.com"]
    assert "sent_at" in parsed.metadata


def test_docx_parser_preserves_paragraph_locations() -> None:
    document = DocxDocument()
    document.add_paragraph("Heading", style="Heading 1")
    document.add_paragraph("Body text")
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)

    parsed = parse_content(
        stream,
        original_name="notes.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert parsed.text == "Heading\n\nBody text"
    assert parsed.metadata["paragraph_count"] == 2
    assert parsed.sections[0].location["style"] == "Heading 1"


def test_empty_and_unsupported_files_are_actionable() -> None:
    with pytest.raises(ParserError, match="no extractable text"):
        parse_content(BytesIO(b"\n\n"), original_name="empty.txt", media_type="text/plain")

    with pytest.raises(UnsupportedFormatError):
        parse_content(
            BytesIO(b"binary"), original_name="archive.exe", media_type="application/octet-stream"
        )
